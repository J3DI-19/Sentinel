from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


path = Path(r"E:\Traceveil\docs\Completion Report.docx")
doc = Document(path)
paragraph = next(p for p in doc.paragraphs if p.text.startswith("Closeout position."))

for run in list(paragraph.runs):
    paragraph._p.remove(run._r)

paragraph.paragraph_format.space_before = Pt(0)
paragraph.paragraph_format.space_after = Pt(6)
paragraph.paragraph_format.line_spacing = 1.08

label = paragraph.add_run("Closeout statement. ")
label.bold = True
label.font.color.rgb = RGBColor(0x1F, 0x4D, 0x78)

body = paragraph.add_run(
    "This report serves as the consolidated project handoff record, bringing together "
    "roadmap reconciliation, validation results, acceptance materials, and a clear account "
    "of team ownership and delivered work. Future roadmap updates can use the evidence "
    "captured here to reconcile outstanding checklist entries while preserving the documented delivery history."
)
body.font.color.rgb = RGBColor(0x24, 0x34, 0x47)

for run in paragraph.runs:
    run.font.name = "Aptos"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(9.5)

doc.save(path)
print(paragraph.text)

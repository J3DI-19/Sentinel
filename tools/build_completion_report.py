from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

OUT = r"E:\Traceveil\docs\Completion Report.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
PALE_BLUE = "E8EEF5"
LIGHT = "F2F4F7"
GREEN = "2E7D32"
PALE_GREEN = "E8F5E9"
AMBER = "7A5A00"
PALE_AMBER = "FFF8E1"
GRAY = "555555"
WHITE = "FFFFFF"
BLACK = "000000"


def set_font(run, size=11, color=BLACK, bold=False, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            cell_margins(cell)


def keep_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, 9, GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_status_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    headers = ["Roadmap area", "Original record", "Current assessment", "Evidence / disposition"]
    for i, value in enumerate(headers):
        c = table.rows[0].cells[i]
        shade(c, NAVY)
        p = c.paragraphs[0]
        r = p.add_run(value)
        set_font(r, 9, WHITE, True)
    repeat_header(table.rows[0])
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        keep_row(table.rows[-1])
        for i, value in enumerate(row):
            if idx % 2:
                shade(cells[i], LIGHT)
            p = cells[i].paragraphs[0]
            r = p.add_run(value)
            set_font(r, 8.5, BLACK, i == 0)
        if row[2] == "Complete":
            shade(cells[2], PALE_GREEN)
            cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(GREEN)
        elif row[2] == "Optional / deferred":
            shade(cells[2], PALE_AMBER)
            cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(AMBER)
    set_table_geometry(table, [1800, 1500, 1600, 4460])
    doc.add_paragraph()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
for level, size, before, after, color in ((1, 16, 16, 8, BLUE), (2, 13, 12, 6, BLUE), (3, 12, 8, 4, NAVY)):
    st = styles[f"Heading {level}"]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
for list_name in ("List Bullet", "List Number"):
    st = styles[list_name]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.paragraph_format.left_indent = Inches(0.5)
    st.paragraph_format.first_line_indent = Inches(-0.25)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.167

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = header.add_run("TRACEVEIL  |  PROJECT COMPLETION REPORT")
set_font(r, 8.5, GRAY, True)
add_page_number(section.footer.paragraphs[0])

# Memo masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(3)
r = p.add_run("COMPLETION REPORT")
set_font(r, 24, NAVY, True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
r = p.add_run("Traceveil roadmap reconciliation, delivery status, and contribution record")
set_font(r, 13, GRAY)
for label, value in (
    ("Prepared", "25 August 2026"),
    ("Source set", "Roadmap, frontend completion audit, Phase 3 demo guide, and contribution record"),
    ("Repository snapshot", "Contribution history reviewed through commit d66a0ac and working-tree integration as of 15 August 2026"),
):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + ": ")
    set_font(r, 10.5, BLACK, True)
    r = p.add_run(value)
    set_font(r, 10.5, BLACK)

add_heading(doc, "Executive summary", 1)
p = doc.add_paragraph()
r = p.add_run("Outcome: ")
set_font(r, 11, GREEN, True)
r = p.add_run("Traceveil's scoped software roadmap is complete.")
set_font(r, 11, BLACK, True)
p.add_run(" The later integration audit confirms all production-route software requirements are implemented, with 43 frontend tests, 60 backend tests, TypeScript/lint, production build, and OpenAPI drift checks passing.")

metrics = doc.add_table(rows=1, cols=3)
metrics.style = "Table Grid"
metric_data = [
    ("38.9%", "Roadmap file recorded", "58 of 149 boxes checked"),
    ("100%", "Scoped software complete", "Per later completion evidence"),
    ("Optional", "Remaining scope", "Physical hardware and MQTT"),
]
for i, (big, label, detail) in enumerate(metric_data):
    c = metrics.rows[0].cells[i]
    shade(c, PALE_BLUE if i != 1 else PALE_GREEN)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(big + "\n")
    set_font(r, 19, GREEN if i == 1 else NAVY, True)
    r = p.add_run(label + "\n")
    set_font(r, 9.5, BLACK, True)
    r = p.add_run(detail)
    set_font(r, 8.5, GRAY)
set_table_geometry(metrics, [3120, 3120, 3120])

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(9)
r = p.add_run("Interpretation. ")
set_font(r, 10, AMBER, True)
r = p.add_run("The 38.9% figure is a documentation snapshot, not the best measure of delivered software. Many unchecked roadmap items were implemented after the checklist was last updated. This report therefore preserves the raw figure and separately reports evidence-adjusted software completion.")
set_font(r, 10, BLACK)

add_heading(doc, "Roadmap completion by area", 1)
status_rows = [
    ("1. Initialization", "Complete", "Complete", "Frontend, backend, SQLite, health checks, and optional Ollama configuration delivered."),
    ("2. Frontend & investigation", "Complete", "Complete", "Application shell, case workspaces, evidence flow, investigation views, assistant, and report review delivered."),
    ("3. Evidence validation", "Complete", "Complete", "Validation, hashing, provenance, errors, live input contracts, and persistence implemented."),
    ("4. Canonical model", "2 items unchecked", "Complete", "Later work added canonical CASAS and TON_IoT telemetry profiles, fixtures, provenance, and normalization."),
    ("5. Live IoT integration", "13 items unchecked", "Complete", "Authenticated HTTP intake, durable sessions/receipts, device state, persistence, analysis handoff, and controlled simulator delivered."),
    ("6. Detection & correlation", "Complete", "Complete", "Deterministic filters, rules, baselines, risk scoring, correlation, timelines, graphs, aggregates, and live alerts delivered."),
    ("7. Real-time delivery", "10 items unchecked", "Complete", "SSE replay/reconnect/deduplication, bounded buffers, live events/alerts/device state, and forensic boundaries delivered."),
    ("8. Investigation APIs", "13 items unchecked", "Complete", "Case, evidence, event, analysis, timeline, graph, analytics, overview, filtering, pagination, and reanalysis APIs delivered."),
    ("9. Visualization + Qwen", "10 items unchecked", "Complete", "Validated safe visualization specifications, grounded optional Ollama assistance, and deterministic fallbacks delivered."),
    ("10. Investigation chat", "9 items unchecked", "Complete", "Persisted case assistant sessions, bounded retrieval, citations, explanations, and offline handling delivered."),
    ("11. Alerts, reports & email", "10 items unchecked", "Complete", "Deterministic PDFs, immutable versions/hashes, approvals, SMTP drafts/sending controls, and audit history delivered."),
    ("12. Live demonstration", "11 items unchecked", "Complete", "The controlled authenticated HTTP simulator is the acceptance path; normal, malformed, reconnect, and suspicious scenarios are documented."),
    ("Physical hardware / MQTT", "Embedded in open items", "Optional / deferred", "Roadmap explicitly treats physical assembly and MQTT as optional future work; neither blocks software acceptance."),
    ("Final validation", "13 items unchecked", "Complete", "Automated verification and demo/operator instructions are recorded; hardware-specific setup documentation remains relevant only if optional hardware is pursued."),
]
add_status_table(doc, status_rows)

add_heading(doc, "Who did what", 1)
add_heading(doc, "J3DI", 2)
add_bullet(doc, "Established the repository, roadmap, React/Vite and FastAPI foundations, SQLite/Ollama configuration, health checks, baseline tests, Traceveil identity, and academic materials.")
add_bullet(doc, "Built the frontend architecture and investigation experience: navigation, cases, evidence import, dashboards, timelines, graphs, analytics, assistant, reports, responsive behavior, accessibility, and branding.")
add_bullet(doc, "Completed the three-phase integration: truthful intake; persisted batch investigations; authenticated live operations; grounded assistance; deterministic reporting; approval-gated delivery; and audit history.")
add_bullet(doc, "Added CASAS and TON_IoT telemetry integration, deterministic OpenAPI/TypeScript boundaries, shared view models, frontend forensic safeguards, release coordination, and final verification documentation.")

add_heading(doc, "Aarya", 2)
add_bullet(doc, "Step 3 evidence validation (commit 66b43fe, 6 August 2026): schemas, profiles, hashing, validation rules, persistence/migrations, dataset fixtures, live telemetry contracts, tests, and documentation.")
add_bullet(doc, "Step 4 canonical normalization (commit 72ad28c, 7 August 2026): versioned canonical events, deterministic IDs, provenance, timestamps, simulation/generic/live/network adapters, integrity checks, tests, and contracts.")
add_bullet(doc, "Step 6 deterministic analytics (commit b4d2337, 11 August 2026): filters, sorting, rules, baselines, bounded risk scoring, correlation, incidents, timelines, graphs, chart aggregates, live alerts, tests, and documentation.")

add_heading(doc, "Ayra", 2)
p = doc.add_paragraph("No unique contribution is recorded in the reviewed Git history. The Ayra branch points to the shared baseline commit b564a3f and contains no commits unique to Ayra.")

add_heading(doc, "Verification and acceptance record", 1)
verification = [
    ("Frontend tests", "43 passed across 10 files"),
    ("Backend/API/persistence tests", "60 passed"),
    ("TypeScript and lint", "Passed with tsc --noEmit"),
    ("Production build", "Passed; Vite transformed 62 modules"),
    ("OpenAPI drift", "Passed"),
    ("Production-route mock audit", "Passed; legacy fixture pages are unreachable from the production router"),
    ("Live acceptance path", "Authenticated HTTP simulator with normal, malformed, reconnect, and suspicious scenarios"),
]
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
for i, h in enumerate(("Gate", "Recorded result")):
    shade(table.rows[0].cells[i], NAVY)
    r = table.rows[0].cells[i].paragraphs[0].add_run(h)
    set_font(r, 9.5, WHITE, True)
repeat_header(table.rows[0])
for idx, (gate, result) in enumerate(verification):
    cells = table.add_row().cells
    keep_row(table.rows[-1])
    if idx % 2:
        shade(cells[0], LIGHT); shade(cells[1], LIGHT)
    r = cells[0].paragraphs[0].add_run(gate); set_font(r, 9, BLACK, True)
    r = cells[1].paragraphs[0].add_run(result); set_font(r, 9, BLACK)
set_table_geometry(table, [2800, 6560])

add_heading(doc, "Closeout note", 1)
p = doc.add_paragraph()
r = p.add_run("Software completion status: COMPLETE. ")
set_font(r, 11, GREEN, True)
r = p.add_run("The roadmap source should be refreshed so its checkboxes match delivered behavior. Physical hardware assembly and MQTT may remain as a separately labeled future enhancement track; they do not change the accepted software completion result.")
set_font(r, 11, BLACK)

add_heading(doc, "Sources reviewed", 2)
for source in (
    "docs/roadmap.md",
    "docs/frontend-checklist-completion.md",
    "docs/phase3-live-demo.md",
    "docs/Contributions.docx",
):
    add_bullet(doc, source)

doc.core_properties.title = "Traceveil Completion Report"
doc.core_properties.subject = "Roadmap completion and contribution attribution"
doc.core_properties.author = "Traceveil Project Team"
doc.core_properties.keywords = "Traceveil, completion, roadmap, contributions"
doc.save(OUT)
print(OUT)

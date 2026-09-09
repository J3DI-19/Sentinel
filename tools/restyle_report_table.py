from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path = r"E:\Traceveil\docs\Completion Report.docx"
doc = Document(path)
table = doc.tables[1]
table.autofit = False

headers = ["ROADMAP AREA", "ROADMAP STATUS", "EVIDENCE STATUS", "DELIVERY EVIDENCE"]
widths = [1750, 1450, 1900, 4260]


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def set_margins(cell, top=110, bottom=110, start=140, end=140):
    props = cell._tc.get_or_add_tcPr()
    margins = props.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for edge, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, size, color, bold=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


for index, cell in enumerate(table.rows[0].cells):
    cell.text = headers[index]
    shade(cell, "0B2545")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_margins(cell, 130, 130, 140, 140)
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            set_font(run, 9, "FFFFFF", True)

for row_index, row in enumerate(table.rows[1:], 1):
    for column_index, cell in enumerate(row.cells):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.0

        if column_index == 0:
            shade(cell, "E8EEF5")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_font(run, 8.8, "0B2545", True)
        elif column_index in (1, 2):
            status = cell.text.strip()
            if status == "Complete":
                fill, color = "E8F5E9", "2E7D32"
            elif "documented" in status.lower() or "partial" in status.lower():
                fill, color = "FFF4DD", "7A5A00"
            else:
                fill, color = "EDF3FA", "1F4D78"
            shade(cell, fill)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_font(run, 8.3, color, True)
        else:
            shade(cell, "F8FAFC" if row_index % 2 == 0 else "FFFFFF")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_font(run, 8.5, "273444")

grid = table._tbl.tblGrid
for child in list(grid):
    grid.remove(child)
for width in widths:
    column = OxmlElement("w:gridCol")
    column.set(qn("w:w"), str(width))
    grid.append(column)

props = table._tbl.tblPr
table_width = props.find(qn("w:tblW"))
table_width.set(qn("w:w"), "9360")
table_width.set(qn("w:type"), "dxa")
indent = props.find(qn("w:tblInd"))
if indent is None:
    indent = OxmlElement("w:tblInd")
    props.append(indent)
indent.set(qn("w:w"), "120")
indent.set(qn("w:type"), "dxa")

for row in table.rows:
    for cell, width in zip(row.cells, widths):
        cell_width = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
        cell_width.set(qn("w:w"), str(width))
        cell_width.set(qn("w:type"), "dxa")
        cell.width = Inches(width / 1440)

borders = props.find(qn("w:tblBorders"))
if borders is None:
    borders = OxmlElement("w:tblBorders")
    props.append(borders)
for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
    node = borders.find(qn(f"w:{edge}"))
    if node is None:
        node = OxmlElement(f"w:{edge}")
        borders.append(node)
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), "4")
    node.set(qn("w:color"), "D7DEE8")

doc.save(path)
print(path)

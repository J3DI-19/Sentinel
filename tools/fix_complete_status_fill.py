from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor


path = Path(r"E:\Traceveil\docs\Completion Report.docx")
document = Document(path)
table = document.tables[1]

updated = 0
for row in table.rows[1:]:
    for column_index in (1, 2):
        cell = row.cells[column_index]
        if cell.text.strip() != "Complete":
            continue

        properties = cell._tc.get_or_add_tcPr()
        for shading in list(properties.findall(qn("w:shd"))):
            properties.remove(shading)

        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), "E8F5E9")
        properties.append(shading)

        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0x2E, 0x7D, 0x32)
        updated += 1

document.save(path)
print(f"Normalized {updated} Complete status cells")

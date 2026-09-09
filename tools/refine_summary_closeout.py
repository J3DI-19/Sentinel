from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


path = Path(r"E:\Traceveil\docs\Completion Report.docx")
doc = Document(path)


def format_run(run, *, bold=False, color="243447"):
    run.font.name = "Aptos"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def rewrite(paragraph, label, text):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    lead = paragraph.add_run(label)
    format_run(lead, bold=True, color="1F4D78")
    body = paragraph.add_run(text)
    format_run(body)


summary_outcome = next(p for p in doc.paragraphs if p.text.startswith("Outcome:"))
summary_interpretation = next(p for p in doc.paragraphs if p.text.startswith("Interpretation."))
closeout = next(p for p in doc.paragraphs if p.text.startswith("Closeout status:"))

rewrite(
    summary_outcome,
    "Delivery position. ",
    "Traceveil's frontend integration checklist is complete, and the reviewed delivery evidence confirms substantial implementation across the core software areas. The roadmap itself currently records 60 of 136 planned items complete (44.1%); that remains the formal whole-roadmap figure until its unchecked items are individually validated and reconciled.",
)

rewrite(
    summary_interpretation,
    "How to read this report. ",
    "A roadmap status of Complete reflects items checked in the source roadmap. An evidence status such as Implemented in later work records delivery supported by later project materials even where the original roadmap boxes remain unchecked. The 100% result applies only to the narrower frontend integration checklist, not to the roadmap as a whole.",
)

rewrite(
    closeout,
    "Closeout position. ",
    "Frontend integration is verified complete within its documented checklist scope. For the overall roadmap, the recorded position remains 60 of 136 items complete (44.1%), while the evidence reviewed in this report demonstrates additional software delivery beyond those checked boxes. The final record therefore distinguishes verified implementation from roadmap reconciliation instead of treating them as the same measure.",
)

doc.save(path)
print("Rewrote executive summary and closeout note")

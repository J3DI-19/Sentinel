from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DOCX = Path(r"E:\Traceveil\docs\Completion Report.docx")
doc = Document(DOCX)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in list(tc_pr.findall(qn("w:shd"))):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=150, start=170, bottom=150, end=170):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = qn(f"w:{edge}")
        node = tc_mar.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "10" if edge in ("top", "bottom") else "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "FFFFFF")
        borders.append(node)
    tbl_pr.append(borders)


def style_run(run, size=9, color="243447", bold=False):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def clear_cell(cell):
    paragraph = cell.paragraphs[0]
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    return paragraph


def add_label(cell, text, color="547087"):
    p = cell.add_paragraph() if cell.paragraphs[0].text else cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text.upper())
    style_run(run, size=7.5, color=color, bold=True)
    return p


def add_body(cell, text, size=8.5, color="243447", bold=False, after=5):
    p = cell.add_paragraph() if cell.paragraphs[0].text else cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    style_run(run, size=size, color=color, bold=bold)
    return p


def add_bullet(cell, text):
    p = cell.add_paragraph(style="List Bullet") if cell.paragraphs[0].text else cell.paragraphs[0]
    if not cell.paragraphs[0].text and p.style.name != "List Bullet":
        p.style = doc.styles["List Bullet"]
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    style_run(run, size=8.25, color="243447")


heading = next(p for p in doc.paragraphs if p.text == "Work allocation")
intro = next(p for p in doc.paragraphs if p.text.startswith("The agreed project responsibilities were allocated"))
verification = next(p for p in doc.paragraphs if p.text == "Verification and acceptance record")

heading.text = "Team ownership & contributions"
heading.style = doc.styles["Heading 1"]
intro.text = (
    "The agreed project responsibilities and reviewed contribution history are combined below. "
    "Allocated scope records assigned ownership; contribution record summarizes the verified work."
)

# Remove the old, duplicated allocation and contribution paragraphs while leaving
# the replacement heading, introduction, and following verification section intact.
body = doc._body._body
children = list(body)
start = children.index(intro._p) + 1
end = children.index(verification._p)
for element in children[start:end]:
    if element.tag == qn("w:p"):
        body.remove(element)

table = doc.add_table(rows=1, cols=3)
widths = [1500, 2460, 5400]
set_table_geometry(table, widths)
set_table_borders(table)

headers = ["TEAM MEMBER", "ALLOCATED SCOPE", "CONTRIBUTION RECORD"]
for index, text in enumerate(headers):
    cell = table.rows[0].cells[index]
    clear_cell(cell)
    set_cell_fill(cell, "18324B")
    set_cell_margins(cell, top=120, bottom=120)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    style_run(run, size=8, color="FFFFFF", bold=True)

records = [
    {
        "name": "J3DI",
        "role": "PROJECT LEAD",
        "accent": "234E70",
        "scope_fill": "EAF2F8",
        "steps": "1  |  2  |  9  |  10",
        "scope": "Documentation; planning; team coordination and leadership; merging; review; integration oversight; and final polish.",
        "contributions": [
            "Served as the overall project lead, owning team coordination, leadership, implementation planning, work sequencing, roadmap management, progress tracking, documentation strategy, integration decisions, and delivery oversight from project initiation through completion.",
            "Established the repository and technical foundation, including the React/Vite frontend, FastAPI backend, SQLite and Ollama configuration, health checks, baseline tests, Traceveil identity, and academic project materials.",
            "Led and built the frontend architecture and investigation experience: navigation, cases, evidence import, dashboards, timelines, graphs, analytics, assistant, reports, responsive behavior, accessibility, and branding.",
            "Coordinated and completed the three-phase integration, dataset-scope corrections, CASAS and TON_IoT telemetry support, OpenAPI and shared-contract boundaries, release integration, branch synchronization, testing and verification, roadmap/status updates, pretesting materials, integration checklists, demo instructions, and final project documentation.",
        ],
    },
    {
        "name": "Aarya",
        "role": "CONTRIBUTOR",
        "accent": "2F6F9F",
        "scope_fill": "EDF5FA",
        "steps": "3  |  4  |  6  |  8",
        "scope": None,
        "contributions": [
            "Step 3 evidence validation (commit 66b43fe, 6 August 2026): schemas, profiles, hashing, validation rules, persistence/migrations, dataset fixtures, live telemetry contracts, tests, and documentation.",
            "Step 4 canonical normalization (commit 72ad28c, 7 August 2026): versioned canonical events, deterministic IDs, provenance, timestamps, simulation/generic/live/network adapters, integrity checks, tests, and contracts.",
            "Step 6 deterministic analytics (commit b4d2337, 11 August 2026): filters, sorting, rules, baselines, bounded risk scoring, correlation, incidents, timelines, graphs, chart aggregates, live alerts, tests, and documentation.",
        ],
    },
    {
        "name": "Ayra",
        "role": "CONTRIBUTOR",
        "accent": "187A7A",
        "scope_fill": "EAF7F5",
        "steps": "5  |  7  |  11  |  12",
        "scope": None,
        "contributions": [
            "No unique contribution is recorded in the reviewed Git history. The Ayra branch points to the shared baseline commit b564a3f and contains no commits unique to Ayra."
        ],
    },
]

for record in records:
    cells = table.add_row().cells
    for cell in cells:
        clear_cell(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    set_cell_fill(cells[0], record["accent"])
    set_cell_fill(cells[1], record["scope_fill"])
    set_cell_fill(cells[2], "F8FAFC")

    name_p = cells[0].paragraphs[0]
    name_p.paragraph_format.space_after = Pt(3)
    name_run = name_p.add_run(record["name"])
    style_run(name_run, size=13, color="FFFFFF", bold=True)
    role_p = cells[0].add_paragraph()
    role_p.paragraph_format.space_after = Pt(0)
    role_run = role_p.add_run(record["role"])
    style_run(role_run, size=7.5, color="DCEAF4", bold=True)

    add_label(cells[1], "Roadmap steps", color=record["accent"])
    steps = add_body(cells[1], record["steps"], size=10, color=record["accent"], bold=True, after=7)
    steps.paragraph_format.keep_with_next = True
    if record["scope"]:
        add_label(cells[1], "Project-wide ownership", color=record["accent"])
        add_body(cells[1], record["scope"], size=8.25, after=0)

    for contribution in record["contributions"]:
        add_bullet(cells[2], contribution)

# Mark header row as repeating and keep each member row intact where possible.
header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
tbl_header = OxmlElement("w:tblHeader")
tbl_header.set(qn("w:val"), "true")
header_tr_pr.append(tbl_header)
for row in table.rows[1:]:
    cant_split = OxmlElement("w:cantSplit")
    row._tr.get_or_add_trPr().append(cant_split)

# Move the completed visual directly after its introduction.
intro._p.addnext(table._tbl)

doc.save(DOCX)
print("Merged allocation and contribution history into a 3-row ownership visual")

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PATH = Path(r"E:\Traceveil\docs\Completion Report.docx")
doc = Document(PATH)


def style_run(run, size, color, bold=False):
    run.font.name = "Aptos"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Aptos")
    fonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def clear_cell(cell):
    p = cell.paragraphs[0]
    for run in list(p.runs):
        p._p.remove(run._r)
    return p


def fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in list(tc_pr.findall(qn("w:shd"))):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def margins(cell, top=155, start=220, bottom=155, end=220):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(cell, color, top="single", bottom="single", left="single", right="single", size=8):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcBorders"))
    if old is not None:
        tc_pr.remove(old)
    container = OxmlElement("w:tcBorders")
    for edge, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), val)
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
        container.append(node)
    tc_pr.append(container)


def set_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_width(cell, widths[index] if len(row.cells) == 2 else sum(widths))


def add_bullet(cell, text):
    p = cell.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    style_run(run, 8.5, "243447")


records = [
    {
        "name": "J3DI",
        "role": "PROJECT LEAD",
        "accent": "234E70",
        "body": "EEF4F8",
        "steps": "01   02   09   10",
        "scope": "Documentation; planning; team coordination and leadership; merging; review; integration oversight; and final polish.",
        "contributions": [
            "Served as the overall project lead, owning team coordination, leadership, implementation planning, work sequencing, roadmap management, progress tracking, documentation strategy, integration decisions, and delivery oversight from project initiation through completion.",
            "Established the repository and technical foundation, including the React/Vite frontend, FastAPI backend, SQLite and Ollama configuration, health checks, baseline tests, Traceveil identity, and academic project materials.",
            "Led and built the frontend architecture and investigation experience: navigation, cases, evidence import, dashboards, timelines, graphs, analytics, assistant, reports, responsive behavior, accessibility, and branding.",
            "Coordinated and completed the three-phase integration, dataset-scope corrections, CASAS and TON_IoT telemetry support, OpenAPI and shared-contract boundaries, release integration, branch synchronization, testing and verification, roadmap/status updates, pretesting materials, integration checklists, demo instructions, and final project documentation.",
        ],
    },
    {
        "name": "Aarya",
        "role": "CONTRIBUTOR",
        "accent": "2F6F9F",
        "body": "F0F6FA",
        "steps": "03   04   06   08",
        "scope": None,
        "contributions": [
            "Step 3 evidence validation (commit 66b43fe, 6 August 2026): schemas, profiles, hashing, validation rules, persistence/migrations, dataset fixtures, live telemetry contracts, tests, and documentation.",
            "Step 4 canonical normalization (commit 72ad28c, 7 August 2026): versioned canonical events, deterministic IDs, provenance, timestamps, simulation/generic/live/network adapters, integrity checks, tests, and contracts.",
            "Step 6 deterministic analytics (commit b4d2337, 11 August 2026): filters, sorting, rules, baselines, bounded risk scoring, correlation, incidents, timelines, graphs, chart aggregates, live alerts, tests, and documentation.",
        ],
    },
    {
        "name": "Ayra",
        "role": "CONTRIBUTOR",
        "accent": "187A7A",
        "body": "EDF8F6",
        "steps": "05   07   11   12",
        "scope": None,
        "contributions": [
            "No unique contribution is recorded in the reviewed Git history. The Ayra branch points to the shared baseline commit b564a3f and contains no commits unique to Ayra."
        ],
    },
]

old_table = next(table for table in doc.tables if table.cell(0, 0).text == "TEAM MEMBER")
anchor = old_table._tbl

table = doc.add_table(rows=0, cols=2)
for record in records:
    header = table.add_row().cells
    body = table.add_row().cells
    body_cell = body[0].merge(body[1])

    for cell in header:
        clear_cell(cell)
        fill(cell, record["accent"])
        margins(cell, top=125, bottom=125)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        borders(cell, record["accent"], size=8)

    name_p = header[0].paragraphs[0]
    name_p.paragraph_format.space_after = Pt(1)
    name = name_p.add_run(record["name"])
    style_run(name, 12.5, "FFFFFF", True)
    role = name_p.add_run(f"   {record['role']}")
    style_run(role, 7.5, "DCEAF4", True)

    steps_p = header[1].paragraphs[0]
    steps_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    steps_p.paragraph_format.space_after = Pt(0)
    label = steps_p.add_run("ASSIGNED STEPS   ")
    style_run(label, 7.5, "DCEAF4", True)
    step_numbers = steps_p.add_run(record["steps"])
    style_run(step_numbers, 10.5, "FFFFFF", True)

    clear_cell(body_cell)
    fill(body_cell, record["body"])
    margins(body_cell, top=165, start=240, bottom=190, end=240)
    body_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    borders(body_cell, record["accent"], top="nil", bottom="single", left="single", right="single", size=8)

    first = body_cell.paragraphs[0]
    if record["scope"]:
        first.paragraph_format.space_after = Pt(3)
        scope_label = first.add_run("ALLOCATED PROJECT-WIDE OWNERSHIP")
        style_run(scope_label, 7.5, record["accent"], True)
        scope_p = body_cell.add_paragraph()
        scope_p.paragraph_format.space_after = Pt(8)
        scope_p.paragraph_format.line_spacing = 1.05
        scope_run = scope_p.add_run(record["scope"])
        style_run(scope_run, 8.75, "243447")
        contribution_label = body_cell.add_paragraph()
    else:
        contribution_label = first
    contribution_label.paragraph_format.space_after = Pt(3)
    label_run = contribution_label.add_run("VERIFIED CONTRIBUTION RECORD")
    style_run(label_run, 7.5, record["accent"], True)
    for contribution in record["contributions"]:
        add_bullet(body_cell, contribution)

    # Keep each header attached to its body while still allowing the content row
    # to flow naturally if Word needs to break it across pages.
    for p in header[0].paragraphs + header[1].paragraphs:
        p.paragraph_format.keep_with_next = True

set_geometry(table, [3500, 5860])

# Replace the previous matrix in place so surrounding report content is untouched.
anchor.addprevious(table._tbl)
anchor.getparent().remove(anchor)

doc.save(PATH)
print("Replaced dense matrix with wide stacked team cards")

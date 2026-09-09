from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


path = Path(r"E:\Traceveil\docs\Completion Report.docx")
doc = Document(path)
table = next(t for t in doc.tables if t.cell(0, 0).text.startswith("J3DI"))

for row_index, row in enumerate(table.rows):
    if row_index % 2 == 1:
        merged_cell = row.cells[0]
        merged_cell._tc.get_or_add_tcPr().tcW.set(qn("w:w"), "9360")
        merged_cell._tc.get_or_add_tcPr().tcW.set(qn("w:type"), "dxa")

doc.save(path)
print("Normalized full-width card rows to 9360 DXA")
